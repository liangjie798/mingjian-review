from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Literal

import fitz
from docx import Document
from openpyxl import load_workbook

from backend.model_provider import ModelConfig, analyze_documents

Severity = Literal["blocking", "high", "medium", "info"]


@dataclass(slots=True)
class Evidence:
    file: str
    page: int
    excerpt: str


@dataclass(slots=True)
class Finding:
    finding_id: str
    severity: Severity
    title: str
    detail: str
    evidence: list[Evidence]
    suggestion: str
    needs_human_review: bool = False
    source: Literal["rule", "model"] = "rule"


@dataclass(slots=True)
class ParsedFile:
    path: Path
    name: str
    pages: int
    characters: int
    parser: str


@dataclass(slots=True)
class ReviewResult:
    files: list[ParsedFile]
    findings: list[Finding]
    model_status: str = "内置规则审查"


SUPPORTED_SUFFIXES = {".pdf", ".docx", ".xlsx", ".txt", ".md", ".csv", ".json"}


def extract_pdf(data: bytes) -> list[str]:
    document = fitz.open(stream=data, filetype="pdf")
    try:
        return [page.get_text("text") for page in document]
    finally:
        document.close()


def extract_docx(data: bytes) -> list[str]:
    document = Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return ["\n".join(lines)]


def extract_xlsx(data: bytes) -> list[str]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        pages: list[str] = []
        for sheet in workbook.worksheets:
            rows = [
                " | ".join("" if value is None else str(value) for value in row)
                for row in sheet.iter_rows(values_only=True)
            ]
            pages.append(f"工作表：{sheet.title}\n" + "\n".join(rows))
        return pages
    finally:
        workbook.close()


def extract_text(path: Path) -> tuple[list[str], str]:
    data = path.read_bytes()
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(data), "PyMuPDF"
    if suffix == ".docx":
        return extract_docx(data), "python-docx"
    if suffix == ".xlsx":
        return extract_xlsx(data), "openpyxl"
    if suffix in {".txt", ".md", ".csv", ".json"}:
        return [data.decode("utf-8", errors="replace")], "文本解析"
    raise ValueError(f"暂不支持 {suffix or '未知'} 格式")


def first_evidence(documents: list[tuple[str, list[str]]], pattern: str) -> Evidence | None:
    regex = re.compile(pattern, re.IGNORECASE)
    for name, pages in documents:
        for page_number, text in enumerate(pages, start=1):
            match = regex.search(text)
            if match:
                start = max(0, match.start() - 40)
                end = min(len(text), match.end() + 70)
                excerpt = re.sub(r"\s+", " ", text[start:end]).strip()
                return Evidence(file=name, page=page_number, excerpt=excerpt)
    return None


def all_text(documents: list[tuple[str, list[str]]]) -> str:
    return "\n".join(page for _, pages in documents for page in pages)


def field_values(documents: list[tuple[str, list[str]]], pattern: str) -> dict[str, Evidence]:
    regex = re.compile(pattern, re.IGNORECASE)
    values: dict[str, Evidence] = {}
    for name, pages in documents:
        for page_number, text in enumerate(pages, start=1):
            for match in regex.finditer(text):
                value = re.sub(r"\s+", " ", match.group(1)).strip(" ：:，,。;；")
                if value:
                    values.setdefault(value, Evidence(name, page_number, match.group(0)[:180]))
    return values


def competition_review(documents: list[tuple[str, list[str]]]) -> list[Finding]:
    findings: list[Finding] = []
    all_names = " ".join(name for name, _ in documents)
    expected = {"报名表": r"报名|申报", "项目书": r"项目书|计划书|商业计划", "承诺书": r"承诺"}
    for label, pattern in expected.items():
        if not re.search(pattern, all_names, re.IGNORECASE):
            findings.append(Finding(
                finding_id=f"C-MISSING-{len(findings) + 1:02d}",
                severity="blocking" if label != "承诺书" else "high",
                title=f"未识别到{label}",
                detail=f"当前文件名中没有识别到{label}，请确认它是否属于必交材料。",
                evidence=[Evidence(file="材料清单", page=1, excerpt=f"待核对材料：{label}")],
                suggestion=f"上传{label}，或在正式规则中确认该材料并非必交。",
                needs_human_review=True,
            ))

    evidence = first_evidence(documents, r"团队成员\s*[:：]\s*([^\n]+)")
    if evidence:
        value = re.split(r"[:：]", evidence.excerpt, maxsplit=1)[-1]
        members = [item.strip() for item in re.split(r"[、,，;；]", value) if item.strip()]
        if 0 < len(members) < 3:
            findings.append(Finding(
                finding_id="C-TEAM-01", severity="blocking", title="团队人数可能不足",
                detail=f"成员字段中识别到约 {len(members)} 名成员，当前演示规则要求 3 至 5 人。",
                evidence=[evidence], suggestion="根据正式比赛通知确认人数要求，并同步更新报名表和项目书。",
                needs_human_review=True,
            ))
    project_names = field_values(documents, r"(?:项目|作品)名称\s*[:：]\s*([^\n|]{2,60})")
    if len(project_names) > 1:
        values = "、".join(list(project_names)[:4])
        evidence = next(iter(project_names.values()))
        findings.append(Finding(
            finding_id="C-CONSIST-01", severity="high", title="项目名称可能不一致",
            detail=f"不同材料中识别到多个项目名称：{values}。名称差异可能导致资格审查退回。",
            evidence=[evidence], suggestion="统一报名表、项目书、承诺书和附件中的项目名称，并核对标点及简称。",
            needs_human_review=True,
        ))
    text = all_text(documents)
    if not re.search(r"1[3-9]\d{9}|(?:电话|手机|联系方式)\s*[:：]", text):
        findings.append(Finding(
            finding_id="C-CONTACT-01", severity="medium", title="未识别到有效联系方式",
            detail="材料正文中没有识别到手机号或明确的联系方式字段。",
            evidence=[Evidence("材料正文", 1, "未检出手机号或联系方式字段")],
            suggestion="在报名表中补充负责人联系方式，并确认号码完整、可用。", needs_human_review=True,
        ))
    commitment_docs = [(name, pages) for name, pages in documents if re.search(r"承诺|声明", name)]
    if commitment_docs and not re.search(r"签字|签名|盖章|公章", all_text(commitment_docs)):
        findings.append(Finding(
            finding_id="C-SIGN-01", severity="high", title="承诺材料可能缺少签署信息",
            detail="已识别到承诺书或声明文件，但可提取文本中没有出现签字或盖章信息。扫描签名仍需人工查看原件。",
            evidence=[Evidence(commitment_docs[0][0], 1, "未检出签字、签名或盖章字段")],
            suggestion="检查承诺书末页是否需要负责人签字、指导教师签字或单位盖章。", needs_human_review=True,
        ))
    if not findings:
        findings.append(Finding(
            finding_id="C-INFO-01", severity="info", title="基础检查未发现阻断项",
            detail="已完成文件名、材料类型和可提取文本检查。语义规则仍需人工确认。",
            evidence=[Evidence(file=documents[0][0], page=1, excerpt="材料已成功解析并进入审查流程。")],
            suggestion="继续核对作品名称、成员、指导教师和签章等跨文档字段。", needs_human_review=True,
        ))
    return findings


def contract_review(documents: list[tuple[str, list[str]]]) -> list[Finding]:
    findings: list[Finding] = []
    payment = first_evidence(documents, r"(?:预付|预付款|首付款)[^。；\n]{0,35}?(\d{1,3})\s*%")
    if payment:
        percentage_match = re.search(r"(\d{1,3})\s*%", payment.excerpt)
        percentage = int(percentage_match.group(1)) if percentage_match else 0
        if percentage > 30:
            findings.append(Finding(
                finding_id="H-PAY-01", severity="high", title="预付款超过规则上限",
                detail=f"识别到预付款比例为 {percentage}%，高于当前内部规则的 30%。",
                evidence=[payment], suggestion="调整付款节点，或提交附带理由的例外审批。",
            ))
    acceptance = first_evidence(documents, r"验收")
    acceptance_period = first_evidence(documents, r"验收[^。；\n]{0,30}\d+\s*(?:个)?(?:工作)?日")
    if acceptance and not acceptance_period:
        findings.append(Finding(
            finding_id="H-ACCEPT-01", severity="medium", title="验收期限未明确",
            detail="材料提到验收，但没有识别到明确的天数或工作日限制。", evidence=[acceptance],
            suggestion="补充验收期限、验收标准和逾期处理方式。", needs_human_review=True,
        ))
    text = all_text(documents)
    if not re.search(r"违约责任|违约金|赔偿责任", text):
        findings.append(Finding(
            finding_id="H-LIABILITY-01", severity="high", title="未识别到违约责任条款",
            detail="合同材料中没有识别到违约责任、违约金或赔偿责任安排。",
            evidence=[Evidence("合同正文", 1, "未检出违约责任相关关键词")],
            suggestion="补充双方违约情形、责任范围、违约金计算方式和损失赔偿规则。", needs_human_review=True,
        ))
    if not re.search(r"争议解决|仲裁|有管辖权的人民法院|诉讼", text):
        findings.append(Finding(
            finding_id="H-DISPUTE-01", severity="high", title="未识别到争议解决条款",
            detail="合同中没有识别到协商、仲裁或诉讼管辖约定。",
            evidence=[Evidence("合同正文", 1, "未检出争议解决、仲裁或法院管辖信息")],
            suggestion="明确争议处理顺序、仲裁机构或有管辖权的法院。", needs_human_review=True,
        ))
    auto_renew = first_evidence(documents, r"(?:自动续期|自动续约|视为续签)[^。；\n]{0,100}")
    if auto_renew:
        findings.append(Finding(
            finding_id="H-RENEW-01", severity="medium", title="存在自动续约安排",
            detail="识别到自动续期或视为续签表述，可能增加退出成本。",
            evidence=[auto_renew], suggestion="明确续约提醒、拒绝续约的通知期限和便捷退出方式。", needs_human_review=True,
        ))
    interpretation = first_evidence(documents, r"最终解释权[^。；\n]{0,50}(?:归|属于)[^。；\n]+")
    if interpretation:
        findings.append(Finding(
            finding_id="H-INTERPRET-01", severity="high", title="存在单方最终解释权表述",
            detail="识别到最终解释权归单方所有的表述，可能造成权利义务失衡。",
            evidence=[interpretation], suggestion="删除单方解释权表述，改为双方协商并依合同目的解释。", needs_human_review=True,
        ))
    if not findings:
        findings.append(Finding(
            finding_id="H-INFO-01", severity="info", title="基础条款检查完成",
            detail="未触发预付款和验收期限规则，仍需人工确认主体、金额和责任条款。",
            evidence=[Evidence(file=documents[0][0], page=1, excerpt="合同材料已成功解析并进入审查流程。")],
            suggestion="结合企业内部制度补充规则后重新运行审查。", needs_human_review=True,
        ))
    return findings


def review_paths(
    scenario: Literal["competition", "contract"],
    paths: list[Path],
    progress: Callable[[int, int, str], None] | None = None,
    model_config: ModelConfig | None = None,
) -> ReviewResult:
    if not paths:
        raise ValueError("请先添加需要审查的材料")
    parsed_files: list[ParsedFile] = []
    documents: list[tuple[str, list[str]]] = []
    total = len(paths)
    for index, path in enumerate(paths, start=1):
        if progress:
            progress(index - 1, total, f"正在解析 {path.name}")
        pages, parser = extract_text(path)
        documents.append((path.name, pages))
        parsed_files.append(ParsedFile(path, path.name, max(1, len(pages)), sum(len(page) for page in pages), parser))
    if progress:
        progress(total, total, "正在执行审查规则")
    findings = contract_review(documents) if scenario == "contract" else competition_review(documents)
    model_status = "内置规则审查"
    if model_config and model_config.enabled:
        if progress:
            progress(total, total, f"正在调用模型 · {model_config.model}")
        try:
            items = analyze_documents(model_config, scenario, documents)
            allowed = {"blocking", "high", "medium", "info"}
            for index, item in enumerate(items, start=1):
                severity = str(item.get("severity", "medium")).lower()
                severity = severity if severity in allowed else "medium"
                try:
                    page = max(1, int(item.get("page", 1)))
                except (TypeError, ValueError):
                    page = 1
                findings.append(Finding(
                    finding_id=f"AI-{index:03d}", severity=severity,
                    title=str(item.get("title", "模型发现待核对问题"))[:120],
                    detail=str(item.get("detail", "请结合原文人工确认。"))[:1000],
                    evidence=[Evidence(str(item.get("file", "模型审查"))[:180], page,
                                       str(item.get("excerpt", "未提供原文摘录"))[:500])],
                    suggestion=str(item.get("suggestion", "请人工复核并补充材料。"))[:800],
                    needs_human_review=True, source="model",
                ))
            model_status = f"内置规则 + {model_config.model} · 新增 {len(items)} 项"
        except RuntimeError as error:
            model_status = f"规则审查完成；模型增强跳过：{error}"
    return ReviewResult(parsed_files, findings, model_status)
