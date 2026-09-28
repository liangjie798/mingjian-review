from __future__ import annotations

import io
import re
import sys
from pathlib import Path
from typing import Literal

import fitz
from docx import Document
from fastapi import FastAPI, File, Form, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel, Field


def resource_path(relative: str) -> Path:
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1]))
    return base / relative


app = FastAPI(title="明鉴审查 API", version="0.2.0", description="多场景材料审查平台的本地接口")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class Evidence(BaseModel):
    file: str
    page: int = Field(ge=1)
    excerpt: str


class Finding(BaseModel):
    finding_id: str
    severity: Literal["blocking", "high", "medium", "info"]
    title: str
    detail: str
    evidence: list[Evidence]
    suggestion: str
    needs_human_review: bool = False


class DemoReviewRequest(BaseModel):
    scenario: Literal["competition", "contract"]


class ParsedFile(BaseModel):
    name: str
    pages: int
    characters: int
    parser: str


class FileReviewResponse(BaseModel):
    files: list[ParsedFile]
    findings: list[Finding]


SCENARIOS = [
    {"id": "competition", "name": "大学生竞赛审查", "status": "ready", "document_types": ["比赛通知", "报名表", "项目书", "证明材料"]},
    {"id": "contract", "name": "合同审查", "status": "ready", "document_types": ["合同正文", "报价单", "交付清单", "内部规则"]},
    {"id": "research", "name": "科研材料审查", "status": "preview", "document_types": ["论文", "配置文件", "结果数据", "README"]},
]


def extract_pdf(data: bytes) -> list[str]:
    document = fitz.open(stream=data, filetype="pdf")
    return [page.get_text("text") for page in document]


def extract_docx(data: bytes) -> list[str]:
    document = Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return ["\n".join(lines)]


def extract_xlsx(data: bytes) -> list[str]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages: list[str] = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(" | ".join("" if value is None else str(value) for value in row))
        pages.append(f"工作表：{sheet.title}\n" + "\n".join(rows))
    return pages


def extract_text(name: str, data: bytes) -> tuple[list[str], str]:
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(data), "PyMuPDF"
    if suffix == ".docx":
        return extract_docx(data), "python-docx"
    if suffix == ".xlsx":
        return extract_xlsx(data), "openpyxl"
    if suffix in {".txt", ".md", ".csv", ".json"}:
        return [data.decode("utf-8", errors="replace")], "UTF-8文本"
    raise ValueError(f"暂不支持 {suffix or '未知'} 格式")


def first_evidence(documents: list[tuple[str, list[str]]], pattern: str) -> Evidence | None:
    regex = re.compile(pattern, re.IGNORECASE)
    for name, pages in documents:
        for page_number, text in enumerate(pages, start=1):
            match = regex.search(text)
            if match:
                start = max(0, match.start() - 35)
                end = min(len(text), match.end() + 55)
                excerpt = re.sub(r"\s+", " ", text[start:end]).strip()
                return Evidence(file=name, page=page_number, excerpt=excerpt)
    return None


def competition_review(documents: list[tuple[str, list[str]]]) -> list[Finding]:
    findings: list[Finding] = []
    all_names = " ".join(name for name, _ in documents)
    expected = {"报名表": r"报名|申报", "项目书": r"项目书|计划书|商业计划", "承诺书": r"承诺"}
    for label, pattern in expected.items():
        if not re.search(pattern, all_names, re.IGNORECASE):
            findings.append(Finding(
                finding_id=f"C-MISSING-{len(findings) + 1:02d}",
                severity="blocking" if label != "承诺书" else "high",
                title=f"未识别到{label}",
                detail=f"当前文件名中没有识别到{label}，请确认是否属于必交材料。",
                evidence=[Evidence(file="材料清单", page=1, excerpt=f"待核对材料：{label}")],
                suggestion=f"上传{label}，或在规则中心将该材料标记为非必交。",
                needs_human_review=True,
            ))

    evidence = first_evidence(documents, r"团队成员\s*[:：]\s*([^\n]+)")
    if evidence:
        value = re.split(r"[:：]", evidence.excerpt, maxsplit=1)[-1]
        members = [item.strip() for item in re.split(r"[、,，;；]", value) if item.strip()]
        if 0 < len(members) < 3:
            findings.append(Finding(
                finding_id="C-TEAM-01", severity="blocking", title="团队人数可能不足",
                detail=f"成员字段中识别到约{len(members)}名成员，演示规则要求3至5人。",
                evidence=[evidence], suggestion="根据正式比赛通知确认人数要求，并同步更新报名表和项目书。",
                needs_human_review=True,
            ))
    if not findings:
        findings.append(Finding(
            finding_id="C-INFO-01", severity="info", title="基础检查未发现阻断项",
            detail="已完成文件名、材料类型和可提取文本检查。语义规则仍需在规则中心确认。",
            evidence=[Evidence(file=documents[0][0], page=1, excerpt="材料已成功解析并进入审查流程。")],
            suggestion="继续核对作品名称、成员、指导教师和签章等跨文档字段。", needs_human_review=True,
        ))
    return findings


def contract_review(documents: list[tuple[str, list[str]]]) -> list[Finding]:
    findings: list[Finding] = []
    payment = first_evidence(documents, r"(?:预付|预付款|首付款)[^。；\n]{0,35}?(\d{1,3})\s*%")
    if payment:
        percentage_match = re.search(r"(\d{1,3})\s*%", payment.excerpt)
        percentage = int(percentage_match.group(1)) if percentage_match else 0
        if percentage > 30:
            findings.append(Finding(
                finding_id="H-PAY-01", severity="high", title="预付款超过演示规则上限",
                detail=f"识别到预付款比例为{percentage}%，高于当前内部规则的30%。",
                evidence=[payment], suggestion="调整付款节点，或提交附带理由的例外审批。",
            ))
    acceptance = first_evidence(documents, r"验收")
    acceptance_period = first_evidence(documents, r"验收[^。；\n]{0,30}\d+\s*(?:个)?(?:工作)?日")
    if acceptance and not acceptance_period:
        findings.append(Finding(
            finding_id="H-ACCEPT-01", severity="medium", title="验收期限未明确",
            detail="材料提到验收，但没有识别到明确的天数或工作日限制。", evidence=[acceptance],
            suggestion="补充验收期限、验收标准和逾期处理方式。", needs_human_review=True,
        ))
    if not findings:
        findings.append(Finding(
            finding_id="H-INFO-01", severity="info", title="基础条款检查完成",
            detail="未触发预付款和验收期限演示规则，仍需人工确认主体、金额和责任条款。",
            evidence=[Evidence(file=documents[0][0], page=1, excerpt="合同材料已成功解析并进入审查流程。")],
            suggestion="添加企业规则包后重新运行审查。", needs_human_review=True,
        ))
    return findings


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"status": "ok", "service": "mingjian-review-api", "version": "0.2.0"}


@app.get("/api/scenarios")
def list_scenarios() -> list[dict[str, object]]:
    return SCENARIOS


@app.post("/api/demo/review", response_model=list[Finding])
def demo_review(payload: DemoReviewRequest) -> list[Finding]:
    if payload.scenario == "contract":
        return contract_review([("设备采购合同.pdf", ["合同约定预付款为80%。设备交付后由甲方组织验收。"])])
    return competition_review([("报名表.pdf", ["团队成员：张同学、李同学"])])


@app.post("/api/review/files", response_model=FileReviewResponse)
async def review_files(
    scenario: Literal["competition", "contract"] = Form(...),
    files: list[UploadFile] = File(...),
) -> FileReviewResponse:
    parsed_files: list[ParsedFile] = []
    documents: list[tuple[str, list[str]]] = []
    for upload in files:
        data = await upload.read()
        pages, parser = extract_text(upload.filename or "未命名文件", data)
        name = upload.filename or "未命名文件"
        documents.append((name, pages))
        parsed_files.append(ParsedFile(
            name=name, pages=max(1, len(pages)), characters=sum(len(page) for page in pages), parser=parser,
        ))
    findings = contract_review(documents) if scenario == "contract" else competition_review(documents)
    return FileReviewResponse(files=parsed_files, findings=findings)


DIST_DIR = resource_path("dist")
if DIST_DIR.exists():
    assets_dir = DIST_DIR / "assets"
    if assets_dir.exists():
        app.mount("/assets", StaticFiles(directory=assets_dir), name="assets")

    @app.get("/{full_path:path}", include_in_schema=False)
    def serve_frontend(full_path: str) -> FileResponse:
        candidate = (DIST_DIR / full_path).resolve()
        if full_path and candidate.is_file() and DIST_DIR.resolve() in candidate.parents:
            return FileResponse(candidate)
        return FileResponse(DIST_DIR / "index.html")
